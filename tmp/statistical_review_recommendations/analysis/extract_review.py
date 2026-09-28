from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document
from lxml import etree


SOURCE = Path(r"C:\Users\Usuario\Downloads\Statistical review — Seed decay temporal weather signatures.docx")
OUT = Path(r"D:\Igor_Masters\Proj_PodAndGrainRot\tmp\statistical_review_recommendations\outputs")
OUT.mkdir(parents=True, exist_ok=True)


def visible_text(element):
    return "".join(element.itertext()).strip()


doc = Document(SOURCE)
lines = []
for idx, paragraph in enumerate(doc.paragraphs, start=1):
    text = paragraph.text.strip()
    if text:
        lines.append(f"P{idx:03d}\t[{paragraph.style.name}]\t{text}")

for table_idx, table in enumerate(doc.tables, start=1):
    lines.append(f"\nTABLE {table_idx}")
    for row_idx, row in enumerate(table.rows, start=1):
        cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
        lines.append(f"R{row_idx:03d}\t" + "\t".join(cells))

(OUT / "review_text.txt").write_text("\n".join(lines), encoding="utf-8")

# Capture text that python-docx may omit because it is stored in tracked-change
# elements, headers, footers, comments, footnotes, or endnotes.
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
xml_lines = []
with ZipFile(SOURCE) as archive:
    candidates = [
        name for name in archive.namelist()
        if name == "word/document.xml"
        or name.startswith("word/header")
        or name.startswith("word/footer")
        or name in {
            "word/comments.xml",
            "word/footnotes.xml",
            "word/endnotes.xml",
        }
    ]
    for name in candidates:
        root = etree.fromstring(archive.read(name))
        xml_lines.append(f"\n## {name}")
        for p_idx, paragraph in enumerate(root.xpath(".//w:p", namespaces=ns), start=1):
            text = "".join(paragraph.xpath(".//w:t/text() | .//w:delText/text()", namespaces=ns)).strip()
            if text:
                xml_lines.append(f"P{p_idx:03d}\t{text}")

(OUT / "review_all_xml_text.txt").write_text("\n".join(xml_lines), encoding="utf-8")
print(f"Extracted {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")
