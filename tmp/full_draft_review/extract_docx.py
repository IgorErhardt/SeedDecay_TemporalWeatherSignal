from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document

src = Path(r"C:\Users\Usuario\Downloads\Modelo_GA_SOJA.docx")
out = Path(r"D:\Igor_Masters\Proj_PodAndGrainRot\tmp\full_draft_review\draft_structure.txt")
doc = Document(src)

lines = []
lines.append(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}")
lines.append("\n=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs, 1):
    text = p.text.replace("\t", "\\t").replace("\n", "\\n")
    if text.strip() or p.style.name.lower().startswith(("heading", "title", "caption")):
        lines.append(f"P{i:04d}\tSTYLE={p.style.name}\t{text}")

lines.append("\n=== TABLES ===")
for ti, table in enumerate(doc.tables, 1):
    lines.append(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows, 1):
        vals = [cell.text.replace("\t", " ").replace("\n", " | ") for cell in row.cells]
        lines.append(f"T{ti}R{ri}\t" + "\t".join(vals))

with ZipFile(src) as zf:
    names = set(zf.namelist())
    media = sorted(n for n in names if n.startswith("word/media/"))
    lines.append("\n=== EMBEDDED MEDIA ===")
    lines.extend(media)
    for part in ("word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"):
        if part not in names:
            continue
        root = etree.fromstring(zf.read(part))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        lines.append(f"\n=== {part} ===")
        for node in root.xpath(".//w:footnote|.//w:endnote|.//w:comment", namespaces=ns):
            text = "".join(node.xpath(".//w:t/text()", namespaces=ns))
            if text.strip():
                lines.append(text)

out.write_text("\n".join(lines), encoding="utf-8")
print(out)
