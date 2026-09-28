from pathlib import Path
from zipfile import ZipFile
import hashlib

from docx import Document

ROOT = Path(r"D:\Igor_Masters\Proj_PodAndGrainRot")
DOCX = ROOT / "tmp" / "full_draft_review" / "Modelo_GA_SOJA_reviewed_current_analysis.docx"

doc = Document(DOCX)
all_text = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    all_text += "\n" + "\n".join("\t".join(c.text for c in r.cells) for r in t.rows)

for forbidden in [
    "seven prespecified 10-day intervals",
    "all seven interval-specific",
    "other six intervals",
    "meteorological-unit aggregation",
    "approximately 72 field trials",
]:
    assert forbidden not in all_text, forbidden

for required in [
    "eight prespecified 10-day intervals",
    "36 ERA5 grid-cell-by-season clusters",
    "models = \"era5\"",
    "sowing-day-plus-cultivar adjustment",
    "cycle-duration-plus-cultivar adjustment",
    "global P = 0.011",
    "global P = 0.083",
]:
    assert required in all_text, required

red_indices = [
    14, 24, 25, 30, 33, 35, 42, 44, 46, 47, 48, 51, 55, 56, 57, 58,
    59, 61, 62, 64, 68, 82, 86, 94, 95, 96, 97, 98, 101, 105, 108,
    110, 111, 112, 113, 114, 117, 119, 121, 123, 124, 130, 131, 132,
    133, 134, 135, 184, 194, 197,
]
for idx in red_indices:
    assert doc.paragraphs[idx].text.strip(), f"empty altered paragraph {idx}"
    assert doc.paragraphs[idx].runs, f"no run in altered paragraph {idx}"
    assert all(r.font.color.rgb is not None and str(r.font.color.rgb) == "FF0000"
               for r in doc.paragraphs[idx].runs if r.text), f"paragraph {idx} not fully red"

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text:
                        assert run.font.color.rgb is not None and str(run.font.color.rgb) == "FF0000"

figure_map = {
    "word/media/image9.png": ROOT / "figures" / "figure1_trial_locations_map.png",
    "word/media/image14.png": ROOT / "figures" / "figure8_ga_distributions.png",
    "word/media/image7.png": ROOT / "figures" / "figure2_all_process_broad_intervals.png",
    "word/media/image5.png": ROOT / "figures" / "figure4_all_process_shared_sensitivities.png",
    "word/media/image13.png": ROOT / "figures" / "figure5_all_process_loso_stability.png",
    "word/media/image11.png": ROOT / "figures" / "figure3_all_process_functional_curves.png",
    "word/media/image4.png": ROOT / "figures" / "figure4b_functional_shared_sensitivities.png",
    "word/media/image3.png": ROOT / "figures" / "figure5b_functional_loso_stability.png",
    "word/media/image2.png": ROOT / "figures" / "figureS3_functional_coefficient_heatmap.png",
    "word/media/image8.png": ROOT / "figures" / "figureS4_broad_residual_vs_fitted.png",
    "word/media/image6.png": ROOT / "figures" / "figureS5_broad_normal_qq.png",
    "word/media/image15.png": ROOT / "figures" / "figureS6_functional_residual_vs_fitted.png",
    "word/media/image12.png": ROOT / "figures" / "figureS7_functional_normal_qq.png",
    "word/media/image10.png": ROOT / "figures" / "figureS8_refund_penalized_curves.png",
    "word/media/image1.png": ROOT / "figures" / "figureS9_refund_penalized_contrasts.png",
}
sha = lambda b: hashlib.sha256(b).hexdigest()
with ZipFile(DOCX) as z:
    assert z.testzip() is None
    for member, source in figure_map.items():
        assert sha(z.read(member)) == sha(source.read_bytes()), member

print(f"PASS: {DOCX}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
