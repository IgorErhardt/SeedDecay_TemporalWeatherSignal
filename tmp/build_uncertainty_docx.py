from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "scalar_on_function_uncertainty_equations.docx"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
MATHML_NS = "http://www.w3.org/1998/Math/MathML"
math_transform = etree.XSLT(etree.parse(str(MML2OMML)))


def set_run_font(run, name="Aptos", size=10.5, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    set_run_font(p.add_run(text))
    return p


def add_equation(doc, mathml):
    """Insert an editable Office Math equation converted from MathML."""
    root = etree.fromstring(mathml.encode("utf-8"))
    converted = math_transform(root).getroot()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p._p.append(converted)
    return p


EQ_BOOTSTRAP_SE = f'''<math xmlns="{MATHML_NS}"><mrow>
  <msub><mi>SE</mi><mrow><mi>boot</mi><mo>,</mo><mi>p</mi></mrow></msub><mo>(</mo><mi>t</mi><mo>)</mo><mo>=</mo>
  <msub><mi>SD</mi><mi>b</mi></msub><mo>{{</mo>
  <msubsup><mover><mi>β</mi><mo>^</mo></mover><mi>p</mi><mrow><mo>*</mo><mo>(</mo><mi>b</mi><mo>)</mo></mrow></msubsup><mo>(</mo><mi>t</mi><mo>)</mo><mo>}}</mo>
</mrow></math>'''

EQ_MAX_DEVIATION = f'''<math xmlns="{MATHML_NS}"><mrow>
  <msub><mi>M</mi><mrow><mi>p</mi><mi>b</mi></mrow></msub><mo>=</mo>
  <munder><mi>max</mi><mrow><mi>t</mi><mo>∈</mo><mi>𝒯</mi></mrow></munder>
  <mo>|</mo><mfrac>
    <mrow><msubsup><mover><mi>β</mi><mo>^</mo></mover><mi>p</mi><mrow><mo>*</mo><mo>(</mo><mi>b</mi><mo>)</mo></mrow></msubsup><mo>(</mo><mi>t</mi><mo>)</mo><mo>−</mo><msub><mover><mi>β</mi><mo>^</mo></mover><mi>p</mi></msub><mo>(</mo><mi>t</mi><mo>)</mo></mrow>
    <mrow><msub><mi>SE</mi><mrow><mi>boot</mi><mo>,</mo><mi>p</mi></mrow></msub><mo>(</mo><mi>t</mi><mo>)</mo></mrow>
  </mfrac><mo>|</mo>
</mrow></math>'''

EQ_CRITICAL = f'''<math xmlns="{MATHML_NS}"><mrow>
  <msubsup><mi>c</mi><mrow><mi>p</mi><mo>,</mo><mn>0.95</mn></mrow><mi>max</mi></msubsup><mo>=</mo>
  <msub><mi>Q</mi><mn>0.95</mn></msub><mo>(</mo><msub><mi>M</mi><mrow><mi>p</mi><mn>1</mn></mrow></msub><mo>,</mo><mo>…</mo><mo>,</mo><msub><mi>M</mi><mrow><mi>p</mi><mi>B</mi></mrow></msub><mo>)</mo>
</mrow></math>'''

EQ_RIBBON = f'''<math xmlns="{MATHML_NS}"><mrow>
  <msub><mover><mi>β</mi><mo>^</mo></mover><mi>p</mi></msub><mo>(</mo><mi>t</mi><mo>)</mo><mo>±</mo>
  <msubsup><mi>c</mi><mrow><mi>p</mi><mo>,</mo><mn>0.95</mn></mrow><mi>max</mi></msubsup>
  <msub><mi>SE</mi><mrow><mi>boot</mi><mo>,</mo><mi>p</mi></mrow></msub><mo>(</mo><mi>t</mi><mo>)</mo>
</mrow></math>'''

EQ_INTEGRATED = f'''<math xmlns="{MATHML_NS}"><mrow>
  <msub><mi>C</mi><mrow><mi>p</mi><mi>k</mi></mrow></msub><mo>=</mo>
  <munder><mo>∑</mo><mrow><mi>t</mi><mo>∈</mo><msub><mi>B</mi><mi>k</mi></msub></mrow></munder>
  <msub><mover><mi>β</mi><mo>^</mo></mover><mi>p</mi></msub><mo>(</mo><mi>t</mi><mo>)</mo>
</mrow></math>'''

EQ_BOOT_INTEGRATED = f'''<math xmlns="{MATHML_NS}"><mrow>
  <msubsup><mi>C</mi><mrow><mi>p</mi><mi>k</mi></mrow><mrow><mo>*</mo><mo>(</mo><mi>b</mi><mo>)</mo></mrow></msubsup><mo>=</mo>
  <munder><mo>∑</mo><mrow><mi>t</mi><mo>∈</mo><msub><mi>B</mi><mi>k</mi></msub></mrow></munder>
  <msubsup><mover><mi>β</mi><mo>^</mo></mover><mi>p</mi><mrow><mo>*</mo><mo>(</mo><mi>b</mi><mo>)</mo></mrow></msubsup><mo>(</mo><mi>t</mi><mo>)</mo>
</mrow></math>'''

EQ_PERCENTILE_CI = f'''<math xmlns="{MATHML_NS}"><mrow><mo>[</mo>
  <msub><mi>Q</mi><mn>0.025</mn></msub><mo>(</mo><msubsup><mi>C</mi><mrow><mi>p</mi><mi>k</mi></mrow><mrow><mo>*</mo><mo>(</mo><mn>1</mn><mo>)</mo></mrow></msubsup><mo>,</mo><mo>…</mo><mo>,</mo><msubsup><mi>C</mi><mrow><mi>p</mi><mi>k</mi></mrow><mrow><mo>*</mo><mo>(</mo><mi>B</mi><mo>)</mo></mrow></msubsup><mo>)</mo><mo>,</mo>
  <msub><mi>Q</mi><mn>0.975</mn></msub><mo>(</mo><msubsup><mi>C</mi><mrow><mi>p</mi><mi>k</mi></mrow><mrow><mo>*</mo><mo>(</mo><mn>1</mn><mo>)</mo></mrow></msubsup><mo>,</mo><mo>…</mo><mo>,</mo><msubsup><mi>C</mi><mrow><mi>p</mi><mi>k</mi></mrow><mrow><mo>*</mo><mo>(</mo><mi>B</mi><mo>)</mo></mrow></msubsup><mo>)</mo>
<mo>]</mo></mrow></math>'''


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)

for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

title_style_ppr = styles["Title"].element.get_or_add_pPr()
for style_border in title_style_ppr.findall(qn("w:pBdr")):
    title_style_ppr.remove(style_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(10)
set_run_font(title.add_run("Scalar on Function Uncertainty Equations"), "Aptos Display", 18, bold=True)
# Remove any theme-provided title border while retaining the built-in Title style.
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

add_body(
    doc,
    "These equations define the 95% simultaneous confidence ribbon for the daily coefficient function and the percentile cluster-bootstrap confidence interval for each prespecified 10-day integrated contrast. The analysis used B = 999 season-stratified meteorological-unit cluster-bootstrap samples."
)

doc.add_heading("Simultaneous confidence ribbon", level=1)
add_body(
    doc,
    "For meteorological process p, beta-hat_p(t) is the coefficient estimated at lag t and beta-hat*_p(t) is the corresponding coefficient from bootstrap sample b. The bootstrap standard error at each lag is the standard deviation of the bootstrap coefficients:"
)
add_equation(doc, EQ_BOOTSTRAP_SE)
add_body(
    doc,
    "For each bootstrap sample, the largest absolute standardized deviation across the complete lag domain T = {-80, ..., -1} is:"
)
add_equation(doc, EQ_MAX_DEVIATION)
add_body(doc, "The simultaneous critical value is the 95th percentile of these bootstrap maximum deviations:")
add_equation(doc, EQ_CRITICAL)
add_body(doc, "The lower and upper limits of the 95% simultaneous ribbon are therefore given by:")
add_equation(doc, EQ_RIBBON)
add_body(
    doc,
    "The critical value is shared across the complete coefficient function, while the bootstrap standard error is calculated separately at each lag. The ribbon therefore changes width across the temporal domain. A region in which the complete ribbon lies above or below zero is distinguishable from zero under the simultaneous whole-curve criterion."
)

integrated_heading = doc.add_heading("Integrated 10 day contrast", level=1)
integrated_heading.paragraph_format.page_break_before = True
add_body(doc, "For prespecified 10-day interval B_k, the integrated contrast is the sum of the estimated daily coefficients:")
add_equation(doc, EQ_INTEGRATED)
add_body(doc, "The corresponding contrast is recalculated in every bootstrap sample:")
add_equation(doc, EQ_BOOT_INTEGRATED)
add_body(doc, "The 95% percentile cluster-bootstrap confidence interval is:")
add_equation(doc, EQ_PERCENTILE_CI)
add_body(
    doc,
    "The integrated contrast represents the estimated cumulative difference in damaged-grain percentage associated with a one-unit higher exposure on every day of the interval, conditional on growing season. If its confidence interval excludes zero, the model supports a nonzero cumulative association over that prespecified interval. This does not imply that every individual daily coefficient differs from zero."
)

doc.add_heading("Interpretive distinction", level=1)
for text in (
    "The simultaneous ribbon evaluates the daily coefficient function while accounting for examination of the complete 80-day domain.",
    "The integrated-contrast interval evaluates the cumulative association within one prespecified 10-day interval.",
    "An integrated contrast can exclude zero even when the daily simultaneous ribbon includes zero because the contrast accumulates information across adjacent days.",
):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(text))

doc.save(OUTPUT)
print(OUTPUT)
